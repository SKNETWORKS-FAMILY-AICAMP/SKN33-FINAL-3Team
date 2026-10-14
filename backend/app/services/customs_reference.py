"""관세청 공식 HSK 엑셀과 원본 CSV에서 기준정보를 준비한다."""
import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

from backend.app.services.customs_stage import ValidationError


def compact_name(name):
    """검색용 주명칭을 추출한다. 법적 범위 확인은 공식 원문으로 한다."""
    text = name
    # 중첩 괄호를 안쪽부터 제거하여 학명·한자·긴 범위 설명을 덜어낸다.
    while re.search(r"\([^()]*\)|\[[^\[\]]*\]", text):
        text = re.sub(r"\([^()]*\)|\[[^\[\]]*\]", "", text)
    text = text.replace("ㆍ", "·").replace("이나 ", "·").replace(" 또는 ", "·")
    text = re.sub(r"그 밖에 이와 유사한", "유사", text)
    text = re.sub(r"그 밖의|그 밖에", "기타", text)
    text = re.sub(r"으로서 따로 분류되지 않은 것|따로 분류되지 않은 것으로 한정한다", "", text)
    text = re.sub(r"\s+", " ", text).strip(" ,·")
    if len(text) > 100:
        head = text[:96]
        cut = max(head.rfind("·"), head.rfind(","), head.rfind(" "))
        text = (head[:cut] if cut > 60 else head) + " 등"
    return text


def summarize_hs_rows(rows):
    """HS6만 같은 판본의 HS4·HS6 원문으로 요약한다. 다른 단위는 원명을 보존한다."""
    parents = {(row[0], row[1]): row[2] for row in rows if len(row[0]) == 4}
    result = []
    for row in sorted(rows, key=lambda row: (len(row[0]), row[0])):
        code, version, display, original, edition = row
        if len(code) == 6:
            parent_raw = parents.get((code[:4], version))
            if not original or not original.strip():
                raise ValidationError(f"HS6 원문 누락: {code}")
            # 원천에 HS4명이 없는 특수 분류는 HS6 원명을 유지하며 상위명을 만들지 않는다.
            parent = compact_name(parent_raw) if parent_raw else ""
            display = compact_name(original)
            # 잔여 분류를 없애는 대신 품목 문맥을 붙인다. '기타'만 표시하지 않는다.
            if parent and re.fullmatch(r"기타(?:의 것)?|그 밖의 것|그 밖에 것", display):
                display = f"{parent}(기타)"
            elif parent and parent not in display:
                display = f"{parent}: {display}"
            if not display:
                raise ValidationError(f"HS 요약명 생성 실패: {code}")
        result.append((code, version, display, original, edition))
    return result


def read_hsk(path, year, hs_ver, source_edition):
    """HSK별 신성질별·성질별 분류 파일의 지정 연도 시트만 읽는다.

    명칭을 추측하지 않는다. 미세분 HS6는 같은 범위인 단일 xxxx0000
    말단명, 또는 해당 HS4의 유일한 xx00 소호 원문으로 펼친다.
    """
    import openpyxl

    if not re.fullmatch(r"[A-Z0-9_]{1,10}", hs_ver):
        raise ValidationError("hs_ver 형식 오류")
    if not source_edition.strip() or len(source_edition) > 55:
        raise ValidationError("source_edition은 공백 제외 1~55자 필요")
    groups = defaultdict(list)
    with Path(path).open("rb") as stream:
        workbook = openpyxl.load_workbook(stream, read_only=True, data_only=True)
        try:
            sheet = workbook[f"{year}년"]
            header = next(sheet.iter_rows(values_only=True))
            if tuple(header[2:6]) != ("세번2단위품명", "세번4단위품명", "세번6단위품명", "세번10단위품명"):
                raise ValidationError("지원하는 관세청 HSK 분류 엑셀 형식이 아님")
            seen = set()
            for row in sheet.iter_rows(min_row=2, values_only=True):
                code = row[1]
                if str(row[0]) != str(year) or not isinstance(code, str) or not re.fullmatch(r"[0-9]{10}", code):
                    raise ValidationError("엑셀 연도/HS10 문자열 검증 실패")
                if code in seen:
                    raise ValidationError(f"엑셀 HS10 중복: {code}")
                seen.add(code)
                names = tuple(x if isinstance(x, str) else "" for x in row[2:6])
                if not names[3].strip():
                    raise ValidationError(f"HS10 명칭 누락: {code}")
                groups[code[:6]].append((code, *names))
        finally:
            workbook.close()
    if not groups:
        raise ValidationError("HS 엑셀에 품목이 없음")

    references = {}
    methods = defaultdict(int)
    unresolved = []
    parent_groups = defaultdict(set)
    for code in groups:
        parent_groups[code[:4]].add(code)

    def put(code, name, method):
        if not name.strip():
            return
        # HS6 표시명만 요약한다. 다른 단위의 원문을 조용히 자르지 않는다.
        if len(code) != 6 and len(name) > 500:
            raise ValidationError(f"HS 원래 품목명이 name_ko의 500자 한도를 초과함: {code}")
        display = name if len(name) <= 500 else name[:499] + "…"
        value = (code, hs_ver, display, name if len(code) == 6 else None,
                 source_edition + ";" + method)
        if code in references and references[code] != value:
            raise ValidationError(f"동일 HS 코드의 명칭 충돌: {code}")
        references[code] = value

    for hs6, rows in groups.items():
        names = {row[3] for row in rows}
        if len(names) != 1:
            raise ValidationError(f"동일 HS6의 공식 명칭 충돌: {hs6}")
        name = next(iter(names))
        method = "HS6"
        if not name.strip():
            if len(rows) == 1 and rows[0][0] == hs6 + "0000":
                name, method = rows[0][4], "UNSPLIT_HS10"
            elif hs6.endswith("00") and parent_groups[hs6[:4]] == {hs6}:
                parent_names = {row[2] for row in rows}
                if len(parent_names) != 1:
                    raise ValidationError(f"미세분 HS4 명칭 충돌: {hs6}")
                name, method = next(iter(parent_names)), "UNSPLIT_HS4"
        if name.strip():
            put(hs6, name, method)
            methods[method] += 1
        else:
            unresolved.append(hs6)
        for code, name2, name4, _, name10 in rows:
            put(code[:2], name2, "HS2")
            put(code[:4], name4, "HS4")
            put(code, name10, "HS10")
    return summarize_hs_rows(list(references.values())), {"hs6_name_methods": dict(methods),
        "hs6_without_hs4": sorted(code for code in references
            if len(code) == 6 and code[:4] not in references),
        "unresolved_hs6": sorted(unresolved), "hs_counts": {
            str(depth): sum(len(code) == depth for code in references) for depth in (2, 4, 6, 10)},
        "workbook_sha256": hashlib.sha256(Path(path).read_bytes()).hexdigest()}


def read_countries(csv_path, known_codes):
    countries = {}
    with Path(csv_path).open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not {"statCd", "statCdCntnKor1", "year"} <= set(reader.fieldnames or []):
            raise ValidationError("국가 사전용 CSV 컬럼 누락")
        for row in reader:
            if row["year"].strip() == "총계":
                continue
            code, name = row["statCd"].strip(), row["statCdCntnKor1"]
            if code not in known_codes or not name.strip() or len(name) > 120:
                raise ValidationError(f"확인되지 않은 국가 코드 또는 명칭: {code}")
            if code in countries and countries[code] != name:
                raise ValidationError(f"동일 국가 코드 명칭 충돌: {code}")
            countries[code] = name
    if not countries:
        raise ValidationError("CSV에 국가가 없음")
    return countries


def seed_references(connection, countries, hs_rows, refresh_names=False):
    """사전 충돌 시 롤백한다. 명시한 경우 같은 원문·출처의 표시명만 갱신한다."""
    try:
        with connection.cursor() as cursor:
            cursor.execute("SELECT option_type,enabled FROM service_option WHERE option_code='CUSTOMS'")
            source = cursor.fetchone()
            if source is None:
                cursor.execute("""INSERT INTO service_option
                    (option_code,option_type,label_ko,value_json) VALUES ('CUSTOMS','SOURCE','관세청',%s)""",
                    (json.dumps({"url": "https://apis.data.go.kr/1220000/nitemtrade/getNitemtradeList"}),))
            elif source != ("SOURCE", 1):
                raise ValidationError("기존 CUSTOMS 출처 설정이 다름")
            for code, name in sorted(countries.items()):
                cursor.execute("SELECT name_ko,country_role FROM country WHERE country_code=%s", (code,))
                old = cursor.fetchone()
                if old is None:
                    cursor.execute("INSERT INTO country (country_code,name_ko,country_role) VALUES (%s,%s,'COUNTRY')", (code, name))
                elif old != (name, "COUNTRY"):
                    raise ValidationError(f"기존 국가 사전 충돌: {code}")
            additions = []
            updates = []
            for row in hs_rows:
                cursor.execute("SELECT hs_code,hs_ver,name_ko,original_name_ko,source_edition FROM hs_code WHERE hs_code=%s AND hs_ver=%s", row[:2])
                old = cursor.fetchone()
                if old is None:
                    additions.append(row)
                elif old != row:
                    if refresh_names and len(row[0]) in (4, 6) and old[3:] == row[3:]:
                        updates.append((row[2], row[0], row[1]))
                    else:
                        raise ValidationError(f"기존 HS 사전 충돌: {row[0]}")
            cursor.executemany("INSERT INTO hs_code (hs_code,hs_ver,name_ko,original_name_ko,source_edition) VALUES (%s,%s,%s,%s,%s)", additions)
            cursor.executemany("UPDATE hs_code SET name_ko=%s WHERE hs_code=%s AND hs_ver=%s", updates)
        connection.commit()
        return {"countries": len(countries), "hs_added": len(additions), "hs_updated": len(updates)}
    except Exception:
        connection.rollback()
        raise
